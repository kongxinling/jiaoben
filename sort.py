#将文件夹中的doc文档另存为docx,并将其按照“非常危险”，“比较危险”，“比较安全”，“非常安全”进行分类
import docx
import os
from win32com import client
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

#将doc文件转换为docx
def doc_to_docx(log,cov):
    pathfile = os.listdir(log)
    for filename in pathfile:
        filedir = log + '/' + filename
        filelist = filename.strip('.').split('.')

        if filelist[4] == "doc":
            #print(filedir)
            ip = '.'.join(filelist[:4])
            #print(ip)
            word = client.Dispatch('Word.Application')
            doc = word.Documents.Open(filedir)  # 目标路径下的文件
            doc.SaveAs(cov + '/' + ip + ".docx", 16)  # 转化后路径下的文件
            doc.Close()
        else:
            continue


# log_d='C:/Users/xinling/Desktop/2'#文件所在的位置
def extract(covert,level1,level2,level3,level4):
    logFiles = os.listdir(covert)
    # 遍历文件夹内所有文件
    for filename in logFiles:
        filepath = covert + '/' + filename
        #print(filepath)
        filelist = filename.strip('.').split('.')
        #print(filename)
        #print(type(filename))
        # 打开文档
        document = Document(filepath)
        tables = document.tables  # 获取文件中的表格集
        table = tables[0]
        count = 0
        if count < 1:
            if "非常危险" in table.cell(0, 1).text:
                count = count + 1
                print(table.cell(0, 1).text)
                newdir = level1 + '/' + filename
                document.save(newdir)
            if "比较危险" in table.cell(0, 1).text:
                count = count + 1
                print(table.cell(0, 1).text)
                newdir = level2 + '/' + filename
                document.save(newdir)

            if "比较安全" in table.cell(0, 1).text:
                count = count + 1
                print(table.cell(0, 1).text)
                newdir = level3 + '/' + filename
                document.save(newdir)

            if "非常安全" in table.cell(0, 1).text:
                count = count + 1
                print(table.cell(0, 1).text)
                newdir = level4 + '/' + filename
                document.save(newdir)
        else:
            continue

log_d=input("请输入doc文件路径：")
covert=input("请输入转换为docx文件的保存路径：")
doc_to_docx(log_d,covert)
#level=input("请输入分类后主机保存路径")
#text=["非常危险","比较危险","比较安全","非常安全"]
level1= input("请输入非常危险的主机保存的路径：")
level2= input("请输入比较危险的主机保存的路径：")
level3= input("请输入比较安全的主机保存的路径：")
level4= input("请输入非常安全的主机保存的路径：")
extract(covert,level1,level2,level3,level4)















